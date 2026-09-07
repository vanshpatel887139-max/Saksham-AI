"""Quiz generation & management routes.

Supports REAL text extraction from uploaded files:
  - TXT   : direct read
  - PDF   : pdfplumber
  - DOCX  : python-docx
  - PPTX  : python-pptx
  - SRT/VTT : transcript parsing (strips timestamps) for video-based learning
"""

import json
import re
import time
import uuid
from io import BytesIO
from fastapi import APIRouter, HTTPException, UploadFile, File, Form
from database import get_db_connection
from models import QuizGenerateRequest, QuizSubmitRequest
from llm import LLMUnavailableError, llm_available, model_name
from routers.logic import generate_quiz_questions, generate_quiz_questions_llm

router = APIRouter(prefix="/api/quiz", tags=["quiz"])


def _generate(content: str, count: int, difficulty: str) -> dict:
    """LLM-first MCQ generation with deterministic fallback."""
    if llm_available():
        try:
            questions = generate_quiz_questions_llm(content, count, difficulty)
            return {
                "questions": questions,
                "generatedBy": f"llm:{model_name()}",
            }
        except (LLMUnavailableError, ValueError):
            pass  # fall back to deterministic generator
    return {
        "questions": generate_quiz_questions(content, count, difficulty),
        "generatedBy": "mock",
    }


def extract_text_from_txt(data: bytes) -> str:
    return data.decode("utf-8", errors="ignore")


def extract_text_from_pdf(data: bytes) -> str:
    try:
        import pdfplumber
    except ImportError as e:
        raise HTTPException(status_code=500, detail="pdfplumber not installed") from e
    parts = []
    with pdfplumber.open(BytesIO(data)) as pdf:
        for page in pdf.pages:
            text = page.extract_text() or ""
            parts.append(text)
    return "\n".join(parts)


def extract_text_from_docx(data: bytes) -> str:
    try:
        import docx
    except ImportError as e:
        raise HTTPException(status_code=500, detail="python-docx not installed") from e
    document = docx.Document(BytesIO(data))
    return "\n".join(p.text for p in document.paragraphs)


def extract_text_from_pptx(data: bytes) -> str:
    try:
        from pptx import Presentation
    except ImportError as e:
        raise HTTPException(status_code=500, detail="python-pptx not installed") from e
    prs = Presentation(BytesIO(data))
    parts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                parts.append(shape.text)
    return "\n".join(parts)


def extract_text_from_transcript(data: bytes, is_vtt: bool) -> str:
    """Parse SRT/VTT subtitle/transcript files, removing timing and numbering."""
    raw = data.decode("utf-8", errors="ignore")
    if is_vtt:
        raw = re.sub(r"^WEBVTT.*?(\n\n|\n)", "", raw, count=1)
    raw = raw.replace("\ufeff", "")
    lines = []
    for line in raw.splitlines():
        line = line.strip()
        if not line:
            continue
        if line.isdigit() and len(line) <= 4:
            continue  # cue number
        if "-->" in line:
            continue  # timestamp
        if re.match(r"^\d{1,2}:\d{2}", line):
            continue
        # inline timestamps (VTT style) inside text
        line = re.sub(r"<[^>]+>", "", line)
        lines.append(line)
    return "\n".join(lines)


def extract_text(filename: str, data: bytes) -> str:
    name = filename.lower().rstrip("?")
    if name.endswith(".pdf"):
        return extract_text_from_pdf(data)
    if name.endswith(".docx"):
        return extract_text_from_docx(data)
    if name.endswith(".pptx") or name.endswith(".ppt"):
        return extract_text_from_pptx(data)
    if name.endswith(".srt"):
        return extract_text_from_transcript(data, is_vtt=False)
    if name.endswith(".vtt"):
        return extract_text_from_transcript(data, is_vtt=True)
    return extract_text_from_txt(data)


@router.post("/generate")
def generate_quiz(req: QuizGenerateRequest):
    """LLM-assisted MCQ generation (falls back to deterministic logic)."""
    if not req.text.strip():
        raise HTTPException(status_code=400, detail="No text provided")
    result = _generate(req.text, req.count, req.difficulty)
    return {
        "quizId": f"quiz-{int(time.time())}",
        "title": f"{req.difficulty} Quiz",
        "questions": result["questions"],
        "generatedBy": result["generatedBy"],
    }


@router.post("/generate/file")
async def generate_from_file(
    file: UploadFile = File(...),
    count: int = Form(5),
    difficulty: str = Form("Medium"),
):
    """Generate a quiz from an uploaded learning material (real text extraction)."""
    data = await file.read()
    if not data:
        raise HTTPException(status_code=400, detail="Empty file")
    filename = file.filename or "material.txt"
    content = extract_text(filename, data)
    if not content.strip():
        raise HTTPException(status_code=400, detail="Could not extract text from file")
    result = _generate(content, count, difficulty)
    return {
        "quizId": f"quiz-{int(time.time())}",
        "title": f"{difficulty} Quiz from {filename}",
        "questions": result["questions"],
        "generatedBy": result["generatedBy"],
    }


@router.post("/submit")
def submit_quiz(req: QuizSubmitRequest):
    conn = get_db_connection()
    try:
        user = conn.execute("SELECT id FROM users WHERE id = ?", (req.user_id,)).fetchone()
        if not user:
            raise HTTPException(status_code=404, detail="User not found")

        score = 0
        for i, q in enumerate(req.questions):
            if i < len(req.answers) and req.answers[i] == q.get("correctAnswer"):
                score += 1

        quiz_id = f"quiz-{uuid.uuid4().hex[:8]}"
        conn.execute(
            "INSERT INTO quizzes (id, user_id, title, date, score, total_questions) VALUES (?, ?, ?, ?, ?, ?)",
            (quiz_id, req.user_id, req.title, time.strftime("%Y-%m-%d"), score, len(req.questions)),
        )
        for i, q in enumerate(req.questions):
            qid = f"qq-{uuid.uuid4().hex[:8]}"
            conn.execute(
                "INSERT INTO quiz_questions (id, quiz_id, question, options, correct_answer, user_answer, explanation, difficulty, source_excerpt) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
                (
                    qid, quiz_id, q.get("question", ""),
                    json.dumps(q.get("options", [])),
                    q.get("correctAnswer", 0),
                    req.answers[i] if i < len(req.answers) else None,
                    q.get("explanation", ""),
                    q.get("difficulty", ""),
                    q.get("sourceExcerpt", ""),
                ),
            )
        conn.commit()
        return {"quizId": quiz_id, "score": score, "total": len(req.questions)}
    finally:
        conn.close()


@router.get("/history/{user_id}")
def quiz_history(user_id: str):
    conn = get_db_connection()
    try:
        rows = conn.execute(
            "SELECT * FROM quizzes WHERE user_id = ? ORDER BY date DESC",
            (user_id,),
        ).fetchall()
        return [
            {
                "id": r["id"],
                "title": r["title"],
                "date": r["date"],
                "score": r["score"],
                "totalQuestions": r["total_questions"],
            }
            for r in rows
        ]
    finally:
        conn.close()